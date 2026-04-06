"""Article text extraction and episode-page resolution using trafilatura."""
import json
import re
from html.parser import HTMLParser
from pathlib import Path
from typing import Optional, Tuple
from urllib.parse import urljoin, urlparse

from trafilatura import extract, fetch_url

_TEXT_LIKE_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".html",
    ".htm",
}
_HTML_SUFFIXES = {".html", ".htm"}

_DOCUMENT_SUFFIXES = {
    ".pdf",
    ".docx",
    ".pptx",
    ".epub",
    ".xlsx",
    ".csv",
}

_DOCUMENT_DEP_HINTS: dict[str, str | None] = {
    ".pdf": "pdfplumber",
    ".docx": "docx",
    ".pptx": "pptx",
    ".epub": "ebooklib",
    ".xlsx": "openpyxl",
    ".csv": None,  # stdlib only
}

_AUDIO_EXTENSIONS = (".mp3", ".m4a", ".wav", ".ogg", ".aac", ".flac")
_INLINE_TRANSCRIPT_MIN_CHARS = 1_500
_TRANSCRIPT_TEXT_HINTS = (" transcript", "transcript ", "full transcript")
_PODCAST_TEXT_HINTS = (
    " podcast",
    "episode ",
    "show notes",
    "listen on",
    "apple podcasts",
    "spotify",
    "overcast",
)
_JSONLD_PODCAST_RE = re.compile(
    r'"@type"\s*:\s*"(?:PodcastEpisode|PodcastSeries|Podcast|AudioObject)"',
    re.IGNORECASE,
)


class _PageSignalsParser(HTMLParser):
    """Collect lightweight media and link signals from one HTML document."""

    def __init__(self):
        super().__init__()
        self.anchors: list[tuple[str, str]] = []
        self.audio_urls: list[str] = []
        self._current_href: str | None = None
        self._current_parts: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        attr_map = {key.lower(): value for key, value in attrs}
        if tag == "a":
            href = attr_map.get("href")
            if href:
                self._current_href = href
                self._current_parts = []
        if tag in {"audio", "source"}:
            src = attr_map.get("src")
            type_attr = (attr_map.get("type") or "").lower()
            if src and (tag == "audio" or type_attr.startswith("audio/") or _looks_like_audio_href(src)):
                self.audio_urls.append(src)

    def handle_data(self, data: str) -> None:
        if self._current_href is not None and data.strip():
            self._current_parts.append(data.strip())

    def handle_endtag(self, tag: str) -> None:
        if tag == "a" and self._current_href is not None:
            self.anchors.append((self._current_href, " ".join(self._current_parts).strip()))
            self._current_href = None
            self._current_parts = []


def _looks_like_audio_href(href: str) -> bool:
    path = urlparse(href).path.lower()
    return any(path.endswith(ext) for ext in _AUDIO_EXTENSIONS)


def _extract_title(downloaded: str, url: str) -> str:
    """Extract a best-effort title from metadata JSON."""
    title = url
    meta_raw = extract(
        downloaded,
        output_format="json",
        with_metadata=True,
        include_comments=False,
        include_tables=False,
    )
    if not meta_raw:
        return title
    try:
        meta = json.loads(meta_raw)
    except (json.JSONDecodeError, AttributeError):
        return title
    return meta.get("title") or url


def _extract_text(downloaded: str) -> str:
    """Extract best-effort article text from downloaded HTML."""
    return extract(downloaded, include_comments=False, include_tables=False) or ""


def _find_transcript_link(parser: _PageSignalsParser, base_url: str) -> str | None:
    """Return the most likely transcript URL linked from the page, if any."""
    for href, text in parser.anchors:
        normalized_text = f" {text.strip().lower()} "
        normalized_href = href.lower()
        if "transcript" in normalized_href or any(hint in normalized_text for hint in _TRANSCRIPT_TEXT_HINTS):
            return urljoin(base_url, href)
    return None


def _find_audio_link(downloaded: str, parser: _PageSignalsParser, base_url: str) -> str | None:
    """Return a best-effort direct audio URL found in the page."""
    for href in parser.audio_urls:
        if _looks_like_audio_href(href):
            return urljoin(base_url, href)
    for href, _text in parser.anchors:
        if _looks_like_audio_href(href):
            return urljoin(base_url, href)
    match = re.search(r'https?://[^"\'>\s]+(?:\.mp3|\.m4a|\.wav|\.ogg|\.aac|\.flac)(?:\?[^"\'>\s]*)?', downloaded, re.IGNORECASE)
    if match:
        return match.group(0)
    return None


def _looks_like_inline_transcript(downloaded: str, text: str) -> bool:
    """Return True when the page appears to contain a real inline transcript."""
    normalized_html = downloaded.lower()
    normalized_text = f" {text.lower()} "
    return (
        len(text) >= _INLINE_TRANSCRIPT_MIN_CHARS
        and (
            "transcript" in normalized_html
            or any(hint in normalized_text for hint in _TRANSCRIPT_TEXT_HINTS)
        )
    )


def _looks_like_podcast_page(downloaded: str, title: str, text: str) -> bool:
    """Return True for likely episode landing pages that should avoid article fallback."""
    normalized_html = downloaded.lower()
    combined_text = f" {title.lower()} {text.lower()} "
    hint_count = sum(1 for hint in _PODCAST_TEXT_HINTS if hint in normalized_html or hint in combined_text)
    if _JSONLD_PODCAST_RE.search(downloaded):
        return True
    return hint_count >= 2


def resolve_article_ingestion(url: str) -> Optional[dict]:
    """
    Resolve a URL into article text, transcript text, direct podcast audio, or explicit podcast failure.

    Returns:
    - {"mode": "article"|"transcript", "text": ..., "title": ..., "metadata_url": ...}
    - {"mode": "podcast_audio", "audio_url": ..., "title": ..., "metadata_url": ...}
    - {"mode": "podcast_unresolved", "title": ..., "reason": ...}
    - None on fetch/extraction failure
    """
    downloaded = fetch_url(url)
    if not downloaded:
        return None

    title = _extract_title(downloaded, url)
    text = _extract_text(downloaded)
    parser = _PageSignalsParser()
    parser.feed(downloaded)

    transcript_url = _find_transcript_link(parser, url)
    if transcript_url:
        transcript = fetch_article(transcript_url)
        if transcript is not None:
            transcript_text, transcript_title = transcript
            return {
                "mode": "transcript",
                "text": transcript_text,
                "title": transcript_title or title,
                "metadata_url": transcript_url,
            }

    if _looks_like_inline_transcript(downloaded, text):
        return {
            "mode": "transcript",
            "text": text,
            "title": title,
            "metadata_url": url,
        }

    audio_url = _find_audio_link(downloaded, parser, url)
    podcast_like = _looks_like_podcast_page(downloaded, title, text)
    if audio_url and podcast_like:
        return {
            "mode": "podcast_audio",
            "audio_url": audio_url,
            "title": title,
            "metadata_url": url,
        }

    if podcast_like:
        reason = (
            f"transcript link could not be ingested for podcast episode page: {url}"
            if transcript_url
            else f"no transcript or audio found for podcast episode page: {url}"
        )
        return {
            "mode": "podcast_unresolved",
            "title": title,
            "reason": reason,
        }

    if not text:
        return None
    return {
        "mode": "article",
        "text": text,
        "title": title,
        "metadata_url": url,
    }


def fetch_article(url: str) -> Optional[Tuple[str, str]]:
    """
    Fetch and extract article text from url.

    Returns (text, title) on success, or None if fetch or extraction fails.
    Never raises — failed URLs should be skipped by the caller, not crash the build.
    """
    downloaded = fetch_url(url)
    if not downloaded:
        return None

    text = _extract_text(downloaded)
    if not text:
        return None

    title = _extract_title(downloaded, url)

    return text, title


def load_text_like_source(path_like: str | Path) -> Tuple[str, str]:
    """Load one narrow local text-like source and return extracted text plus title."""
    path = Path(path_like).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"Missing local source: {path}")
    if not path.is_file():
        raise ValueError(f"Unsupported local source type: {path}")

    suffix = path.suffix.lower()
    if suffix not in _TEXT_LIKE_SUFFIXES:
        raise ValueError(f"Unsupported local source type: {path}")

    raw_text = path.read_text(encoding="utf-8")
    text = raw_text
    if suffix in _HTML_SUFFIXES:
        text = extract(raw_text, include_comments=False, include_tables=False) or raw_text

    normalized = text.strip()
    if not normalized:
        raise ValueError(f"Local source is empty: {path}")
    return normalized, path.stem


def load_document_source(path_like: str | Path) -> Tuple[str, str]:
    """Load a document source (.pdf, .docx, .pptx, .epub, .xlsx, .csv) and return (text, title)."""
    path = Path(path_like).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"Missing local source: {path}")
    if not path.is_file():
        raise ValueError(f"Not a file: {path}")

    suffix = path.suffix.lower()
    if suffix not in _DOCUMENT_SUFFIXES:
        raise ValueError(f"Unsupported document type: {path}")

    module_name = _DOCUMENT_DEP_HINTS.get(suffix)
    if module_name is not None:
        try:
            __import__(module_name)
        except ImportError:
            raise ImportError(
                f"{suffix} support requires '{module_name.split('.')[0]}'. "
                "Install the matching Python package in the same environment as beyin, "
                "then rerun the build. Run `beyin check-deps` for details."
            )

    text = ""
    if suffix == ".pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        text = "\n\n".join(p for p in pages if p.strip())

    elif suffix == ".docx":
        import docx
        doc = docx.Document(str(path))
        text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

    elif suffix == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(path))
        slides = []
        for slide in prs.slides:
            parts = [
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            ]
            if parts:
                slides.append("\n".join(parts))
        text = "\n\n".join(slides)

    elif suffix == ".epub":
        import ebooklib
        from ebooklib import epub
        from html.parser import HTMLParser

        class _StripHTML(HTMLParser):
            def __init__(self):
                super().__init__()
                self._parts: list[str] = []
            def handle_data(self, data: str) -> None:
                self._parts.append(data)
            def get_text(self) -> str:
                return " ".join(self._parts)

        book = epub.read_epub(str(path))
        chapters = []
        for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
            parser = _StripHTML()
            parser.feed(item.get_content().decode("utf-8", errors="ignore"))
            chapter_text = parser.get_text().strip()
            if chapter_text:
                chapters.append(chapter_text)
        text = "\n\n".join(chapters)

    elif suffix == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        rows = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join(str(cell) for cell in row if cell is not None)
                if row_text.strip():
                    rows.append(row_text)
        text = "\n".join(rows)

    elif suffix == ".csv":
        import csv
        rows = []
        with open(path, newline="", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                row_text = "\t".join(cell for cell in row if cell.strip())
                if row_text.strip():
                    rows.append(row_text)
        text = "\n".join(rows)

    normalized = text.strip()
    if not normalized:
        raise ValueError(f"Local source is empty or could not be extracted: {path}")
    return normalized, path.stem
