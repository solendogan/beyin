"""Article text extraction using trafilatura."""
import json
from pathlib import Path
from typing import Optional, Tuple

from trafilatura import extract, fetch_url

_TEXT_LIKE_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".rst",
    ".html",
    ".htm",
}
_HTML_SUFFIXES = {".html", ".htm"}

_DOCUMENT_SUFFIXES = {
    ".pdf",
    ".docx",
    ".pptx",
    ".epub",
    ".xlsx",
    ".csv",
}

_DOCUMENT_DEP_HINTS: dict[str, tuple[str, str]] = {
    ".pdf": ("pdfplumber", "pip install pdfplumber"),
    ".docx": ("docx", "pip install python-docx"),
    ".pptx": ("pptx", "pip install python-pptx"),
    ".epub": ("ebooklib", "pip install ebooklib"),
    ".xlsx": ("openpyxl", "pip install openpyxl"),
    ".csv": (None, None),  # stdlib only
}


def fetch_article(url: str) -> Optional[Tuple[str, str]]:
    """
    Fetch and extract article text from url.

    Returns (text, title) on success, or None if fetch or extraction fails.
    Never raises — failed URLs should be skipped by the caller, not crash the build.
    """
    downloaded = fetch_url(url)
    if not downloaded:
        return None

    text = extract(downloaded, include_comments=False, include_tables=False)
    if not text:
        return None

    # Extract title from metadata JSON (second extract call with output_format="json")
    title = url  # fallback to URL if metadata unavailable
    meta_raw = extract(
        downloaded,
        output_format="json",
        with_metadata=True,
        include_comments=False,
        include_tables=False,
    )
    if meta_raw:
        try:
            meta = json.loads(meta_raw)
            title = meta.get("title") or url
        except (json.JSONDecodeError, AttributeError):
            pass  # keep URL as fallback title

    return text, title


def load_text_like_source(path_like: str | Path) -> Tuple[str, str]:
    """Load one narrow local text-like source and return extracted text plus title."""
    path = Path(path_like).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"Missing local source: {path}")
    if not path.is_file():
        raise ValueError(f"Unsupported local source type: {path}")

    suffix = path.suffix.lower()
    if suffix not in _TEXT_LIKE_SUFFIXES:
        raise ValueError(f"Unsupported local source type: {path}")

    raw_text = path.read_text(encoding="utf-8")
    text = raw_text
    if suffix in _HTML_SUFFIXES:
        text = extract(raw_text, include_comments=False, include_tables=False) or raw_text

    normalized = text.strip()
    if not normalized:
        raise ValueError(f"Local source is empty: {path}")
    return normalized, path.stem


def load_document_source(path_like: str | Path) -> Tuple[str, str]:
    """Load a document source (.pdf, .docx, .pptx, .epub, .xlsx, .csv) and return (text, title)."""
    path = Path(path_like).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"Missing local source: {path}")
    if not path.is_file():
        raise ValueError(f"Not a file: {path}")

    suffix = path.suffix.lower()
    if suffix not in _DOCUMENT_SUFFIXES:
        raise ValueError(f"Unsupported document type: {path}")

    module_name, install_hint = _DOCUMENT_DEP_HINTS.get(suffix, (None, None))
    if module_name is not None:
        try:
            __import__(module_name)
        except ImportError:
            raise ImportError(
                f"{suffix} support requires '{module_name.split('.')[0]}'. "
                f"Install it with: {install_hint}"
            )

    text = ""
    if suffix == ".pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        text = "\n\n".join(p for p in pages if p.strip())

    elif suffix == ".docx":
        import docx
        doc = docx.Document(str(path))
        text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

    elif suffix == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(path))
        slides = []
        for slide in prs.slides:
            parts = [
                shape.text.strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and shape.text.strip()
            ]
            if parts:
                slides.append("\n".join(parts))
        text = "\n\n".join(slides)

    elif suffix == ".epub":
        import ebooklib
        from ebooklib import epub
        from html.parser import HTMLParser

        class _StripHTML(HTMLParser):
            def __init__(self):
                super().__init__()
                self._parts: list[str] = []
            def handle_data(self, data: str) -> None:
                self._parts.append(data)
            def get_text(self) -> str:
                return " ".join(self._parts)

        book = epub.read_epub(str(path))
        chapters = []
        for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
            parser = _StripHTML()
            parser.feed(item.get_content().decode("utf-8", errors="ignore"))
            chapter_text = parser.get_text().strip()
            if chapter_text:
                chapters.append(chapter_text)
        text = "\n\n".join(chapters)

    elif suffix == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        rows = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join(str(cell) for cell in row if cell is not None)
                if row_text.strip():
                    rows.append(row_text)
        text = "\n".join(rows)

    elif suffix == ".csv":
        import csv
        rows = []
        with open(path, newline="", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                row_text = "\t".join(cell for cell in row if cell.strip())
                if row_text.strip():
                    rows.append(row_text)
        text = "\n".join(rows)

    normalized = text.strip()
    if not normalized:
        raise ValueError(f"Local source is empty or could not be extracted: {path}")
    return normalized, path.stem
